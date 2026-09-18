"""Profile UCI Online Retail II

Writes source-derived daily workload statistics. These statistics describe the
historical UCI records only; they are not actual workload at a modeled site.
"""

from __future__ import annotations

import csv
import argparse
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_FILE = PROJECT_ROOT / "data/raw/uci/online_retail_II.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "data/processed/uci_daily_workload_profile.csv"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-sheet", help="Profile only this workbook sheet")
    parser.add_argument("--output", type=Path, help="Override the output CSV path")
    args = parser.parse_args()
    output_file = args.output or OUTPUT_FILE

    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    daily: dict[tuple[str, str], dict[str, object]] = defaultdict(
        lambda: {
            "source_lines": 0,
            "positive_quantity_lines": 0,
            "positive_units": 0,
            "nonpositive_quantity_lines": 0,
            "invoice_ids": set(),
        }
    )

    for worksheet in workbook.worksheets:
        if args.source_sheet and worksheet.title != args.source_sheet:
            continue
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows)
        columns = {name: index for index, name in enumerate(header)}

        for row in rows:
            invoice_date = row[columns["InvoiceDate"]]
            quantity = row[columns["Quantity"]]
            invoice = row[columns["Invoice"]]
            if invoice_date is None:
                continue

            key = (worksheet.title, invoice_date.date().isoformat())
            record = daily[key]
            record["source_lines"] += 1
            record["invoice_ids"].add(str(invoice))

            if quantity is not None and quantity > 0:
                record["positive_quantity_lines"] += 1
                record["positive_units"] += quantity
            else:
                record["nonpositive_quantity_lines"] += 1

    with output_file.open("w", newline="") as destination:
        writer = csv.DictWriter(
            destination,
            fieldnames=[
                "source_sheet",
                "planning_date",
                "source_lines",
                "distinct_invoice_ids",
                "positive_quantity_lines",
                "positive_units",
                "nonpositive_quantity_lines",
            ],
        )
        writer.writeheader()
        for (sheet, date), record in sorted(daily.items()):
            writer.writerow(
                {
                    "source_sheet": sheet,
                    "planning_date": date,
                    "source_lines": record["source_lines"],
                    "distinct_invoice_ids": len(record["invoice_ids"]),
                    "positive_quantity_lines": record["positive_quantity_lines"],
                    "positive_units": record["positive_units"],
                    "nonpositive_quantity_lines": record["nonpositive_quantity_lines"],
                }
            )

    print(f"Wrote {len(daily)} daily source rows to {output_file}")


if __name__ == "__main__":
    main()
