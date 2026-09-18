"""Create exact-row signatures for the December 2010 UCI overlap audit."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_FILE = PROJECT_ROOT / "data/raw/uci/online_retail_II.xlsx"
OVERLAP_START = "2010-12-01"
OVERLAP_END = "2010-12-09"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-sheet", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    worksheet = workbook[args.source_sheet]
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows)
    columns = {name: index for index, name in enumerate(header)}
    signature_columns = [
        "Invoice", "StockCode", "Description", "Quantity", "InvoiceDate",
        "Price", "Customer ID", "Country",
    ]

    written = 0
    with args.output.open("w", newline="") as destination:
        writer = csv.DictWriter(destination, fieldnames=["signature", *signature_columns])
        writer.writeheader()
        for row in rows:
            invoice_date = row[columns["InvoiceDate"]]
            if invoice_date is None:
                continue
            date = invoice_date.date().isoformat()
            if not OVERLAP_START <= date <= OVERLAP_END:
                continue
            values = [row[columns[column]] for column in signature_columns]
            normalized = [value.isoformat() if hasattr(value, "isoformat") else str(value) for value in values]
            writer.writerow(dict(zip(["signature", *signature_columns], ["|".join(normalized), *normalized])))
            written += 1

    print(f"Wrote {written} overlap rows from {args.source_sheet} to {args.output}")


if __name__ == "__main__":
    main()
