"""Create a privacy-safe aggregate profile of the Soto-Pinedo workbooks."""
from __future__ import annotations

import json
import math
import statistics
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data/raw/Soto-Pinedo DataSet"
OUTPUT = ROOT / "data/processed/soto_pinedo_profile.json"


def workbook(prefix: str) -> Path:
    matches = sorted(SOURCE.glob(f"{prefix}*.xlsx"))
    if len(matches) != 1:
        raise RuntimeError(f"Expected one workbook matching {prefix!r}; found {len(matches)}")
    return matches[0]


def percentile(values: list[float], probability: float) -> float:
    ordered = sorted(values)
    if not ordered:
        raise ValueError("Cannot calculate a percentile for an empty collection")
    position = (len(ordered) - 1) * probability
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return ordered[lower]
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)


def numeric_summary(values) -> dict:
    clean = [float(value) for value in values if isinstance(value, (int, float)) and not isinstance(value, bool)]
    if not clean:
        return {"observations": 0}
    return {
        "observations": len(clean),
        "mean": round(statistics.fmean(clean), 3),
        "median": round(statistics.median(clean), 3),
        "p95": round(percentile(clean, .95), 3),
        "minimum": round(min(clean), 3),
        "maximum": round(max(clean), 3),
    }


def matrix_values(sheet) -> list[float]:
    return [value for row in sheet.iter_rows(min_row=2, min_col=2, values_only=True)
            for value in row if isinstance(value, (int, float)) and not isinstance(value, bool)]


def parse_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def picking_profile() -> dict:
    book = load_workbook(workbook("DATA PICKING"), read_only=True, data_only=True)
    try:
        return {
            "daily_picking_errors": numeric_summary(matrix_values(book["ErroresDiarios"])),
            "daily_travel_distance_m": numeric_summary(matrix_values(book["DistanciaDiaria"])),
            "order_preparation_time_min": numeric_summary(matrix_values(book["TiempoPorPedido"])),
            "orders_per_day": numeric_summary(matrix_values(book["PedidosDia"])),
            "operator_columns": 8,
            "classification": "External benchmark; observed-versus-simulated provenance is not established in the archive",
        }
    finally:
        book.close()


def preparation_2025_profile() -> dict:
    book = load_workbook(workbook("DATA ML 2025"), read_only=True, data_only=True)
    try:
        sheet = book["DATA 2025"]
        rows = sheet.iter_rows(values_only=True)
        header = [str(value).strip() if value is not None else "" for value in next(rows)]
        index = {name: position for position, name in enumerate(header)}
        start_column = index["FECHA-HORA DE LANZAMIENTO"]
        end_column = index["FECHA-HORA TERMINO DE PREPARACIÓN"]
        status_column = index["STATUS"]
        durations = []
        statuses = Counter()
        covered_rows = 0
        total_rows = 0
        date_values = []
        for row in rows:
            total_rows += 1
            statuses[str(row[status_column]).strip() if row[status_column] is not None else "Missing"] += 1
            start = parse_datetime(row[start_column])
            end = parse_datetime(row[end_column])
            if start:
                date_values.append(start.date())
            if start and end:
                hours = (end - start).total_seconds() / 3600
                if 0 <= hours <= 24 * 7:
                    durations.append(hours)
                    covered_rows += 1
        return {
            "rows": total_rows,
            "valid_preparation_intervals": covered_rows,
            "interval_coverage_share": round(covered_rows / total_rows, 4) if total_rows else 0,
            "preparation_cycle_hours": numeric_summary(durations),
            "launch_date_start": min(date_values).isoformat() if date_values else None,
            "launch_date_end": max(date_values).isoformat() if date_values else None,
            "status_counts": dict(sorted(statuses.items())),
            "privacy": "Only aggregate durations and status counts retained; names, customers, carriers, plates, and IDs excluded",
        }
    finally:
        book.close()


def fulfillment_profile() -> dict:
    book = load_workbook(workbook("DATA IDENTIFIC"), read_only=True, data_only=True)
    try:
        sheet = book["DATA"]
        rows = sheet.iter_rows(values_only=True)
        header = [str(value).strip() if value is not None else "" for value in next(rows)]
        index = {name: position for position, name in enumerate(header)}
        requested_column = index["Solicitado"]
        fulfilled_column = index["Fill Rate"]
        reason_column = index["Motivo de Fill Rate"]
        requested = fulfilled = 0.0
        reasons = Counter()
        records = 0
        for row in rows:
            records += 1
            if isinstance(row[requested_column], (int, float)):
                requested += float(row[requested_column])
            if isinstance(row[fulfilled_column], (int, float)):
                fulfilled += float(row[fulfilled_column])
            reason = row[reason_column]
            if reason is not None:
                reasons[str(reason).strip()] += 1
        return {
            "records": records,
            "requested_quantity": round(requested, 3),
            "fulfilled_quantity": round(fulfilled, 3),
            "weighted_fill_share": round(fulfilled / requested, 4) if requested else None,
            "top_fulfillment_reasons": [{"reason": reason, "records": count}
                                        for reason, count in reasons.most_common(8)],
            "privacy": "Customer, employee, location-detail, and order identifiers excluded",
        }
    finally:
        book.close()


def build() -> dict:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Missing extracted source folder: {SOURCE}")
    return {
        "metadata": {
            "title": "Soto-Pinedo privacy-safe operational benchmark profile",
            "source": "User-supplied Soto-Pinedo workbook archive",
            "license_status": "Not established; no license file was included in the archive",
            "allowed_use": "Internal research profile only until license and provenance are confirmed",
            "geographic_scope": "Peru; external to the U.S. network capacity model",
            "created_at": datetime.now(timezone.utc).isoformat(),
        },
        "picking_benchmarks": picking_profile(),
        "preparation_2025": preparation_2025_profile(),
        "fulfillment_coverage": fulfillment_profile(),
        "excluded_fields": [
            "customer names and customer IDs", "vehicle plates", "carrier names",
            "employee or programmer identifiers", "order and shipment identifiers",
        ],
        "model_use": "Descriptive external benchmark only; values do not replace U.S. site productivity or staffing assumptions",
    }


if __name__ == "__main__":
    result = build()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(result, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    print(f"Wrote {OUTPUT.relative_to(ROOT)}")
