"""Extract selected role evidence from the pinned August 2026 O*NET release."""
import csv
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
ROLES={'53-7065.00':'fulfillment_associate','53-7064.00':'packing_associate',
       '43-5071.00':'shipping_receiving_coordinator','53-1042.00':'warehouse_supervisor'}


def main():
    source=ROOT/'data/raw/db_31_0_excel'
    readme=(source/'Read Me.txt').read_text()
    if 'O*NET 31.0 Database' not in readme or 'August 2026 Release' not in readme:
        raise ValueError('Expected O*NET 31.0 August 2026 source')
    out=ROOT/'tableau/onet_2026'
    out.mkdir(parents=True,exist_ok=True)
    manifest=dict(database_version='31.0',release='2026-08',
                  attribution='O*NET 31.0 Database, U.S. Department of Labor, Employment and Training Administration',
                  license='https://creativecommons.org/licenses/by/4.0/',
                  source='https://www.onetcenter.org/database.html',
                  transformation='Selected four occupations; skill importance scale only, no ranking or productivity inference.',
                  limits='Record dates may predate the release. Not staffing, wages, training duration or worker qualifications.',
                  tables={},source_hashes={'Read Me.txt':hashlib.sha256((source/'Read Me.txt').read_bytes()).hexdigest()})
    for name in ['Occupation Data','Task Statements','Essential Skills','Transferable Skills']:
        path=source/(name+'.xlsx')
        workbook=load_workbook(path,read_only=True,data_only=True)
        iterator=workbook.active.iter_rows(values_only=True)
        headers=next(iterator)
        rows=[]
        for values in iterator:
            row=dict(zip(headers,values))
            code=row['O*NET-SOC Code']
            if code not in ROLES or ('Scale ID' in row and row['Scale ID']!='IM'):
                continue
            rows.append(dict(role_id=ROLES[code],database_version='31.0',release_month='2026-08',**row))
        workbook.close()
        assert {r['O*NET-SOC Code'] for r in rows}==set(ROLES), name
        filename=name.lower().replace(' ','_')+'.csv'
        with (out/filename).open('w',newline='') as handle:
            writer=csv.DictWriter(handle,fieldnames=list(rows[0]))
            writer.writeheader();writer.writerows(rows)
        manifest['source_hashes'][path.name]=hashlib.sha256(path.read_bytes()).hexdigest()
        manifest['tables'][filename]=dict(rows=len(rows),record_dates_by_role={role:sorted({r['Date'] for r in rows if r['role_id']==role and r.get('Date')}) for role in ROLES.values()})
    (out/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print(json.dumps(manifest['tables'],indent=2))


if __name__=='__main__':
    main()
